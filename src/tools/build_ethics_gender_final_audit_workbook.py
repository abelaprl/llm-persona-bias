"""Build one consolidated Excel workbook for the final ETHICS gender audit."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[2]
V15 = ROOT / "result" / "ETHICS" / "reports" / "Gender_Persona" / "linguistic_markers_gender_candidate_20260821_v1_5_0" / "audit_codex_r1_20260822_v6_final"
V16 = ROOT / "result" / "ETHICS" / "reports" / "Gender_Persona" / "linguistic_markers_gender_candidate_20260821_v1_6_0"
REGRESSION = ROOT / "docs" / "audit_artifacts_ethics_gender_v16_20260822_v8_final"
OUTPUT = V16 / "audit_codex_r1_20260822_final"
WORKBOOK = OUTPUT / "audit_ethics_gender_codex_r1_final.xlsx"


def main() -> None:
    if OUTPUT.exists() and any(OUTPUT.iterdir()):
        raise FileExistsError(f"Refusing to overwrite non-empty output: {OUTPUT}")
    OUTPUT.mkdir(parents=True, exist_ok=True)

    v15_manifest = json.loads((V15 / "audit_manifest.json").read_text(encoding="utf-8"))
    v16_manifest = json.loads((V16 / "manifest.json").read_text(encoding="utf-8"))
    regression_manifest = json.loads((REGRESSION / "manifest.json").read_text(encoding="utf-8"))
    retained_precision = (
        regression_manifest["reviewed_true_positives_retained"]
        / (
            regression_manifest["reviewed_true_positives_retained"]
            + regression_manifest["reviewed_false_positives"]
            - regression_manifest["reviewed_false_positives_eliminated"]
        )
    )
    tp_retention = (
        regression_manifest["reviewed_true_positives_retained"]
        / regression_manifest["reviewed_true_positives"]
    )
    fp_elimination = (
        regression_manifest["reviewed_false_positives_eliminated"]
        / regression_manifest["reviewed_false_positives"]
    )

    overview = pd.DataFrame(
        [
            {"metric": "status", "value": "candidate_pending_independent_human_audit"},
            {"metric": "reviewer", "value": "Codex_R1 (single AI reviewer)"},
            {"metric": "v16_pipeline_version", "value": v16_manifest["pipeline_version"]},
            {"metric": "v16_lexicon_version", "value": v16_manifest["lexicon_version"]},
            {"metric": "full_dataset_rows", "value": v16_manifest["rows_input"]},
            {"metric": "text_available", "value": v16_manifest["responses_text_available"]},
            {"metric": "abstain_included", "value": v16_manifest["responses_abstain"]},
            {"metric": "missing", "value": v16_manifest["responses_missing"]},
            {"metric": "markers", "value": v16_manifest["marker_count"]},
            {"metric": "v15_precision_rows", "value": v15_manifest["precision_rows"]},
            {"metric": "v15_precision_micro", "value": v15_manifest["precision_micro"]},
            {"metric": "v15_precision_macro", "value": v15_manifest["precision_macro"]},
            {"metric": "v15_recall_rows", "value": v15_manifest["recall_rows"]},
            {"metric": "v15_recall_false_negatives", "value": v15_manifest["recall_false_negatives"]},
            {"metric": "v15_reviewed_away_hits", "value": v15_manifest["recall_reviewed_away_automatic_hits"]},
            {"metric": "v16_reviewed_fp_eliminated", "value": regression_manifest["reviewed_false_positives_eliminated"]},
            {"metric": "v16_fp_elimination_rate", "value": fp_elimination},
            {"metric": "v16_reviewed_tp_retained", "value": regression_manifest["reviewed_true_positives_retained"]},
            {"metric": "v16_tp_retention_rate", "value": tp_retention},
            {"metric": "v16_retained_span_precision", "value": retained_precision},
            {"metric": "v16_recall_false_negatives", "value": regression_manifest["v16_recall_false_negatives"]},
            {"metric": "v16_recall_false_positives", "value": regression_manifest["v16_recall_false_positives"]},
            {"metric": "claim_limit", "value": "Regression evidence only; not independently human validated."},
        ]
    )
    guide = pd.DataFrame(
        [
            {"step": 1, "sheet": "v15_precision", "instruction": "Filter is_true_positive=0 and judge matched_text together with match_context."},
            {"step": 2, "sheet": "v15_precision_sum", "instruction": "Markers below precision 0.90 fail the current promotion threshold."},
            {"step": 3, "sheet": "v15_recall_diff", "instruction": "Inspect every false-negative and reviewed-away automatic marker at full-response level."},
            {"step": 4, "sheet": "v16_reg_precision", "instruction": "Check which v1.5 TP/FP spans are retained or eliminated by v1.6."},
            {"step": 5, "sheet": "v16_reg_recall", "instruction": "Confirm response-level v1.6 marker sets against the corrected manual sets."},
            {"step": 6, "sheet": "v16_open_forms", "instruction": "Confirm every observed open gender form is registered in the closed lexicon."},
            {"step": 7, "sheet": "v16_unmatched", "instruction": "Review unmatched candidate contexts; generic mentions are not automatically persona uptake."},
            {"step": 8, "sheet": "audit_corrections", "instruction": "Review all contextual relabeling made during audit passes 5-6."},
            {"step": 9, "sheet": "README", "instruction": "Do not promote from candidate until independent human review and agreement are complete."},
        ]
    )
    tables = {
        "README": overview,
        "HOW_TO_AUDIT": guide,
        "v15_precision_sum": pd.read_csv(V15 / "precision_summary_codex_r1.csv"),
        "v15_precision": pd.read_csv(V15 / "audit_marker_hits_codex_r1.csv"),
        "v15_recall_sum": pd.read_csv(V15 / "recall_summary_codex_r1.csv"),
        "v15_recall_diff": pd.read_csv(V15 / "recall_discrepancies_codex_r1.csv"),
        "v15_recall": pd.read_csv(V15 / "audit_recall_responses_codex_r1.csv"),
        "audit_corrections": pd.read_csv(V15 / "audit_correction_log.csv"),
        "v16_reg_prec_sum": pd.read_csv(REGRESSION / "precision_regression_summary.csv"),
        "v16_reg_precision": pd.read_csv(REGRESSION / "precision_regression_rows.csv"),
        "v16_reg_rec_sum": pd.read_csv(REGRESSION / "recall_regression_summary.csv"),
        "v16_reg_recall": pd.read_csv(REGRESSION / "recall_regression_rows.csv"),
        "v16_gender_coverage": pd.read_csv(V16 / "gender_candidate_context_coverage.csv"),
        "v16_open_forms": pd.read_csv(V16 / "open_gender_forms.csv"),
        "v16_unmatched": pd.read_csv(V16 / "gender_unmatched_context_audit.csv"),
    }
    with pd.ExcelWriter(WORKBOOK, engine="openpyxl") as writer:
        for sheet_name, table in tables.items():
            table.to_excel(writer, sheet_name=sheet_name, index=False)
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.row_dimensions[1].height = 28
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for column in sheet.columns:
                letter = column[0].column_letter
                sample = list(column[:250])
                width = min(70, max(12, max(len(str(cell.value or "")) for cell in sample) + 2))
                sheet.column_dimensions[letter].width = width
                header = str(column[0].value or "")
                if any(token in header for token in ("context", "visible_text", "note", "instruction", "reason")):
                    for cell in column[1:]:
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
            if sheet.title == "v15_precision_sum":
                precision_column = next(
                    (cell.column for cell in sheet[1] if cell.value == "precision"), None
                )
                if precision_column:
                    for row in range(2, sheet.max_row + 1):
                        cell = sheet.cell(row=row, column=precision_column)
                        if isinstance(cell.value, (float, int)) and cell.value < 0.9:
                            cell.fill = PatternFill("solid", fgColor="F4CCCC")
            if sheet.title == "v16_reg_precision":
                outcome_column = next(
                    (cell.column for cell in sheet[1] if cell.value == "v16_regression_outcome"), None
                )
                if outcome_column:
                    for row in range(2, sheet.max_row + 1):
                        cell = sheet.cell(row=row, column=outcome_column)
                        if cell.value == "false_positive_eliminated":
                            cell.fill = PatternFill("solid", fgColor="D9EAD3")
                        elif cell.value in {"true_positive_lost", "false_positive_retained"}:
                            cell.fill = PatternFill("solid", fgColor="F4CCCC")

    manifest = {
        "status": "candidate_pending_independent_human_audit",
        "workbook": str(WORKBOOK),
        "v15_audit": str(V15),
        "v16_full_run": str(V16),
        "v16_regression": str(REGRESSION),
        "sheet_rows": {name: len(table) for name, table in tables.items()},
    }
    (OUTPUT / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
